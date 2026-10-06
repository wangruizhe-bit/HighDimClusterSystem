"""阶段五：实际文件导出、格式验证、公式注入防护与数据不可变性。"""
import numpy as np
import pandas as pd
import pytest
from matplotlib.figure import Figure
from openpyxl import load_workbook

from core.reporting import (
    ResultItem, build_catalog, export_csv, export_excel, export_figure,
    export_html_report, _sheet_name,
)
from core.pca import fit_pca
from core.dimension_reduction import fit_reduction
from core.clustering import fit_clustering, search_k


@pytest.fixture
def data():
    rng = np.random.default_rng(34)
    z = rng.normal(size=(38, 4))
    return pd.DataFrame(z, columns=list("ABCD")).assign(label="kind")


@pytest.fixture
def catalog(data):
    pca = fit_pca(data, selection="manual", component_count=3)
    reduction = fit_reduction(data, method="factor", n_components=2)
    cluster = fit_clustering(data, method="kmeans", n_clusters=3)
    ksearch = search_k(data, k_min=2, k_max=4)
    stats = ("样本统计", data[["A", "B"]].describe().reset_index())
    return build_catalog(data, data_name="analysis.csv", pca=pca,
                         reductions=[reduction], clusters=[cluster],
                         k_search=ksearch, feature_result=stats)


def test_catalog_only_real_present_results(data, catalog):
    minimal = build_catalog(data, data_name="test")
    assert [r.key for r in minimal] == ["working_data", "data_summary"]
    assert len(catalog) > 13
    assert len({item.key for item in catalog}) == len(catalog)
    assert {item.category for item in catalog} >= {"数据集", "PCA", "降维", "聚类", "K 值辅助选择"}
    assert catalog[0].frame.equals(data)
    catalog[0].frame.loc[0, "A"] = 8888
    assert data.loc[0, "A"] != 8888


def test_csv_roundtrip_with_chinese_columns(tmp_path, catalog):
    item = catalog[0]
    path = tmp_path / "真实数据.csv"
    export_csv(item, path)
    restored = pd.read_csv(path)
    assert restored.shape == item.frame.shape
    np.testing.assert_allclose(restored["A"].to_numpy(), item.frame["A"].to_numpy())
    assert restored["label"].tolist() == item.frame["label"].tolist()


def test_excel_all_sheets_are_readable(tmp_path, catalog):
    path = tmp_path / "分析结果.xlsx"
    export_excel(catalog, path)
    wb = load_workbook(path, read_only=True, data_only=False)
    assert len(wb.sheetnames) == len(catalog)
    assert len({name.lower() for name in wb.sheetnames}) == len(catalog)
    assert all(len(name) <= 31 for name in wb.sheetnames)
    wb.close()
    one = pd.read_excel(path, sheet_name=0)
    assert len(one) == len(catalog[0].frame)
    assert "A" in one.columns


def test_excel_sheet_duplicate_names_and_unsafe_user_strings(tmp_path):
    frame = pd.DataFrame({
        "User": ["=HYPERLINK(\"https://malicious.invalid\")", "+SUM(1,1)", "-2+3",
                 "@cmd", "ordinary"],
        "value": [-1.0, 2.2, 3.3, 4.4, 5.5],
    })
    original = frame.copy(deep=True)
    items = [ResultItem(f"key{i}", "相同标题" * 8, "检查", frame) for i in range(3)]
    output = tmp_path / "safe.xlsx"
    export_excel(items, output)
    wb = load_workbook(output, data_only=False)
    assert len(set(wb.sheetnames)) == 3
    for name in wb.sheetnames:
        sheet = wb[name]
        assert sheet["A2"].data_type == "s"
        assert sheet["A2"].value.startswith("'=")
        assert sheet["A3"].value.startswith("'+")
        assert sheet["B2"].data_type == "n"
        assert sheet["B2"].value == -1
    wb.close()
    path = tmp_path / "safe.csv"
    export_csv(items[0], path)
    assert b"'=HYPERLINK" in path.read_bytes()
    pd.testing.assert_frame_equal(frame, original)


def test_html_local_offline_and_escaped(tmp_path):
    table = pd.DataFrame({"x": ["<script>alert('x')</script>", "plain"],
                          "score": [1.23, 9.5]})
    file = tmp_path / "report.html"
    export_html_report([ResultItem("s", "<恶意标题>", "测试", table)], file)
    text = file.read_text(encoding="utf-8")
    assert "<!DOCTYPE html>" in text
    assert "&lt;script&gt;" in text
    assert "<script>" not in text
    assert "&lt;恶意标题&gt;" in text
    assert "1.23" in text
    assert "http://" not in text and "https://" not in text


def test_export_charts_png_and_svg(tmp_path):
    fig = Figure()
    ax = fig.add_subplot(111)
    ax.plot([1, 2, 3], [3, 2, 1])
    png = tmp_path / "plot.png"
    svg = tmp_path / "plot.svg"
    export_figure(fig, png)
    export_figure(fig, svg)
    assert png.read_bytes().startswith(b"\x89PNG")
    assert "<svg" in svg.read_text(encoding="utf-8")


@pytest.mark.parametrize("suffix", [".txt", ".xls", ".json"])
def test_wrong_file_extension_rejected(tmp_path, catalog, suffix):
    with pytest.raises(ValueError):
        export_csv(catalog[0], tmp_path / ("wrong" + suffix))
    with pytest.raises(ValueError):
        export_excel(catalog, tmp_path / ("wrong" + suffix))
    with pytest.raises(ValueError):
        export_html_report(catalog, tmp_path / ("wrong" + suffix))


def test_empty_reports_and_empty_figures_rejected(tmp_path):
    with pytest.raises(ValueError):
        export_excel([], tmp_path / "empty.xlsx")
    with pytest.raises(ValueError):
        export_html_report([], tmp_path / "empty.html")
    with pytest.raises(ValueError):
        export_figure(Figure(), tmp_path / "empty.png")


def test_excel_sheet_name_rules():
    names = set()
    first = _sheet_name(" a:?*[]\\longlonglonglonglonglonglonglong", names)
    second = _sheet_name(" a:?*[]\\longlonglonglonglonglonglonglong", names)
    assert first != second
    assert len(first) <= 31 and len(second) <= 31
    assert not any(c in first for c in "[]:*?/\\")
